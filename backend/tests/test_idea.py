"""Step 3.1: research plan schema, Idea agent (FakeLLM), context files, workspace API."""

import copy
import io
import uuid
from typing import Any

import pytest
from httpx import AsyncClient
from pydantic import ValidationError
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.agents import registry
from app.agents.idea.agent import MissingIdeaError, make_idea_agent, plan_from_state
from app.agents.idea.schemas import IdeaPlan
from app.core.config import get_settings
from app.db.models import AgentRun, Feedback, ProjectMember, User
from app.db.models.enums import MemberRole, RunStatus
from app.services.context import PER_FILE_CHARS, extract_text, format_context, load_context
from app.services.llm import LLMService
from app.services.llm_providers import FakeLLM
from tests.conftest import TEST_STORAGE
from tests.helpers import IDEA, Flow, auth_headers, factory_for

PLAN: dict[str, Any] = {
    "refined_idea": "Rank drug candidates with graph neural networks trained on binding-affinity data.",
    "domain": "Computational drug discovery",
    "keywords": ["graph neural networks", "drug discovery", "binding affinity"],
    "subtopics": ["molecular graph representation", "affinity prediction"],
    "problem_statement": "Screening candidates experimentally is slow and expensive.",
    "motivation": "Better early ranking could cut wet-lab costs.",
    "scope": "In scope: small molecules and public affinity data. Out of scope: wet-lab validation.",
    "expected_contributions": ["A GNN ranking model", "A benchmark comparison protocol"],
    "objectives": [
        {"id": "O1", "text": "Design a GNN that predicts binding affinity."},
        {"id": "O2", "text": "Evaluate ranking quality against baselines."},
        {"id": "O3", "text": "Analyse failure cases by molecular family."},
    ],
    "research_questions": [
        {"id": "RQ1", "text": "Does the GNN rank candidates better than fingerprints?", "kind": "question",
         "linked_objective_ids": ["O1", "O2"]},
        {"id": "RQ2", "text": "Ranking errors concentrate in specific molecular families.", "kind": "hypothesis",
         "linked_objective_ids": ["O3"]},
    ],
    "research_directions": ["Transfer to new targets", "Uncertainty-aware ranking"],
    "title_options": [
        "Graph Neural Networks for Ranking Drug Candidates",
        "GNN-Rank: Affinity-Driven Candidate Prioritisation",
        "Can Graph Neural Networks Rank Drug Candidates Better?",
    ],
}  # fmt: skip


def plan(**changes: Any) -> dict[str, Any]:
    data = copy.deepcopy(PLAN)
    data.update(changes)
    return data


# ─── Plan rules ────────────────────────────────────────────────────────


def test_valid_plan() -> None:
    p = IdeaPlan.model_validate(PLAN)
    assert [o.id for o in p.objectives] == ["O1", "O2", "O3"]
    assert p.research_questions[1].kind == "hypothesis"


@pytest.mark.parametrize(
    ("changes", "message"),
    [
        ({"research_questions": [{**PLAN["research_questions"][0], "linked_objective_ids": ["O9"]},
                                 PLAN["research_questions"][1]]}, "links to objectives that don't exist: O9"),
        ({"research_questions": [PLAN["research_questions"][1],
                                 {**PLAN["research_questions"][0], "id": "RQ3", "linked_objective_ids": ["O3"]}]},
         "none for O1, O2"),
        ({"objectives": [*PLAN["objectives"][:2], {"id": "O2", "text": "Duplicate."}]}, "ids must be unique"),
        ({"objectives": [{"id": "Objective-1", "text": "x"}, *PLAN["objectives"][1:]]}, 'look like "O1"'),
        ({"research_questions": [{**PLAN["research_questions"][0], "linked_objective_ids": []},
                                 PLAN["research_questions"][1]]}, "at least 1 item"),
        ({"objectives": PLAN["objectives"][:2]}, "at least 3 items"),
        ({"research_questions": [PLAN["research_questions"][0],
                                 {**PLAN["research_questions"][1], "text": "Do errors cluster by family?"}]},
         "RQ2 is marked as a hypothesis but is phrased as a question"),
        ({"title_options": PLAN["title_options"][:2]}, "at least 3 items"),
    ],
)  # fmt: skip
def test_plan_rules(changes: dict[str, Any], message: str) -> None:
    with pytest.raises(ValidationError, match=message):
        IdeaPlan.model_validate(plan(**changes))


# ─── Agent ────────────────────────────────────────────────────────────


def llm_with(db_session: AsyncSession, responses: list[Any]) -> tuple[FakeLLM, Any]:
    fake = FakeLLM(list(responses))
    settings = get_settings().model_copy(update={"llm_model_fallback": "", "llm_max_retries": 0})
    return fake, lambda: LLMService(fake, settings, factory_for(db_session))


def state(**extra: Any) -> dict[str, Any]:
    pid = str(extra.pop("project_id", uuid.uuid4()))
    return {
        "project_id": pid,
        "idea": {"raw_idea": IDEA, "domain": "Bioinformatics"},
        "venue": {"name": "ieee"},
        **extra,
    }


async def test_agent_builds_plan_from_idea_and_notes(db_session: AsyncSession, project_id: uuid.UUID) -> None:
    fake, llm = llm_with(db_session, [PLAN])

    async def notes(_: uuid.UUID) -> str:
        return '<file name="meeting.md">\nSupervisor: focus on kinase targets.\n</file>'

    update = await make_idea_agent(llm, notes)(state(project_id=project_id))

    prompt = fake.calls[0].prompt
    assert IDEA in prompt and "Bioinformatics" in prompt and "ieee" in prompt
    assert "Supervisor: focus on kinase targets." in prompt
    assert "Previous plan" not in prompt  # first run
    assert fake.calls[0].system is not None and "Do not cite papers" in fake.calls[0].system
    assert fake.calls[0].json_schema is not None and fake.calls[0].json_schema["title"] == "IdeaPlan"

    assert update["idea"]["raw_idea"] == IDEA
    assert update["idea"]["title_options"] == PLAN["title_options"]
    assert update["title"] == PLAN["title_options"][0]  # defaults to the first option
    assert [o["id"] for o in update["objectives"]] == ["O1", "O2", "O3"]
    assert update["research_questions"][0]["linked_objective_ids"] == ["O1", "O2"]
    assert update["problem"]["statement"] == PLAN["problem_statement"]
    output = update["stage_outputs"]["idea"]
    assert output["agent"] == "idea.plan" and output["section_draft"] is None
    assert "### Objectives" in output["draft"] and "**O1.**" in output["draft"]

    run = await db_session.get(AgentRun, uuid.UUID(output["agent_run_id"]))
    assert run is not None and run.status == RunStatus.SUCCESS and run.agent == "idea.plan"
    assert run.output is not None and run.output["prompt"] == "idea.plan@2"


async def test_regenerate_sends_previous_plan_and_instructions(db_session: AsyncSession, project_id: uuid.UUID) -> None:
    previous = await make_idea_agent(llm_with(db_session, [PLAN])[1], _no_notes)(state(project_id=project_id))
    fake, llm = llm_with(db_session, [plan(domain="Medicinal chemistry")])

    await make_idea_agent(llm, _no_notes)(
        state(project_id=project_id, **previous, pending_feedback="Add a fairness objective")
    )

    prompt = fake.calls[0].prompt
    assert "Add a fairness objective" in prompt
    assert "Previous plan" in prompt and "Design a GNN that predicts binding affinity." in prompt


async def test_invalid_links_are_repaired_by_the_gateway(db_session: AsyncSession, project_id: uuid.UUID) -> None:
    broken = plan(research_questions=[{**PLAN["research_questions"][0], "linked_objective_ids": ["O7"]},
                                      PLAN["research_questions"][1]])  # fmt: skip
    fake, llm = llm_with(db_session, [broken, PLAN])

    update = await make_idea_agent(llm, _no_notes)(state(project_id=project_id))

    assert len(fake.calls) == 2
    assert "links to objectives that don't exist: O7" in fake.calls[1].prompt
    assert update["research_questions"][0]["linked_objective_ids"] == ["O1", "O2"]


async def test_agent_needs_a_raw_idea(db_session: AsyncSession, project_id: uuid.UUID) -> None:
    with pytest.raises(MissingIdeaError):
        await make_idea_agent(llm_with(db_session, [])[1], _no_notes)(
            state(project_id=project_id, idea={"raw_idea": "  "})
        )


def test_plan_round_trips_through_state() -> None:
    from app.agents.idea.agent import plan_to_state

    p = IdeaPlan.model_validate({**PLAN, "title": PLAN["title_options"][1]})
    assert plan_from_state({"project_id": "x", **plan_to_state(p, IDEA)}) == p
    assert plan_from_state({"idea": {"raw_idea": IDEA}}) is None


async def _no_notes(_: uuid.UUID) -> str:
    return ""


# ─── Context files ─────────────────────────────────────────────────────


def test_extract_text_from_txt_pdf_docx() -> None:
    import docx
    import pymupdf

    assert extract_text(b"plain notes", "text/plain") == "plain notes"

    pdf = pymupdf.open()
    pdf.new_page().insert_text((72, 72), "PDF notes about GNNs")
    assert "PDF notes about GNNs" in extract_text(pdf.tobytes(), "application/pdf")

    document = docx.Document()
    document.add_paragraph("Word notes about kinases")
    buf = io.BytesIO()
    document.save(buf)
    mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    assert "Word notes about kinases" in extract_text(buf.getvalue(), mime)

    assert extract_text(b"\x89PNG", "image/png") == ""


async def test_load_context_reads_notes_within_budget(flow: Flow, db_session: AsyncSession, storage: Any) -> None:
    async def upload(name: str, data: bytes, kind: str) -> None:
        resp = await flow.client.post(
            flow.url("/files"), files={"file": (name, data, "text/plain")}, data={"kind": kind}, headers=flow.headers
        )
        assert resp.status_code == 201, resp.text

    await upload("meeting.md", b"Focus on kinase inhibitors.", "notes")
    await upload("long.txt", b"x" * (PER_FILE_CHARS + 500), "notes")
    await upload("results.csv", b"a,b\n1,2\n", "result")  # not a notes file

    docs = await load_context(db_session, TEST_STORAGE, flow.pid)

    assert [d.filename for d in docs] == ["meeting.md", "long.txt"]
    assert docs[1].truncated and len(docs[1].text) == PER_FILE_CHARS
    text = format_context(docs)
    assert '<file name="meeting.md">' in text and '<file name="long.txt" (excerpt)>' in text


# ─── Workspace API (agents backed by a FakeLLM) ──────────────────────────


@pytest.fixture
def idea_llm(monkeypatch: pytest.MonkeyPatch, db_session: AsyncSession) -> FakeLLM:
    """Install the real Idea agent, backed by a FakeLLM and the test database/storage."""
    fake = FakeLLM()
    settings = get_settings().model_copy(update={"llm_model_fallback": "", "llm_max_retries": 0})

    def agents(session_factory: Any) -> dict[Any, Any]:
        async def notes(pid: uuid.UUID) -> str:
            return format_context(await load_context(db_session, TEST_STORAGE, pid))

        from app.db.models.enums import StageName

        return {StageName.IDEA: make_idea_agent(lambda: LLMService(fake, settings, session_factory), notes)}

    monkeypatch.setattr(registry, "agent_factory", agents)
    return fake


async def test_generate_edit_and_approve_plan(flow: Flow, idea_llm: FakeLLM, db_session: AsyncSession) -> None:
    """Acceptance: idea → coherent plan → edit an objective → approve."""
    idea_llm.responses.append(PLAN)
    started = await flow.post("/stages/idea/run", {"instructions": "Keep it to one paper's worth of work"})
    assert started.status_code == 202, started.text
    await flow.finish_jobs()
    assert "Keep it to one paper's worth of work" in idea_llm.calls[0].prompt

    ws = (await flow.client.get(flow.url("/idea"), headers=flow.headers)).json()
    assert ws["status"] == "awaiting_review" and ws["editable"] is True
    assert ws["plan"]["objectives"][0]["text"] == PLAN["objectives"][0]["text"]
    assert ws["plan"]["title"] == PLAN["title_options"][0] and ws["agent_run_id"]
    # The plan isn't written into the Abstract section (that's Step 3.3)
    assert (await flow.client.get(flow.url("/sections/abstract"), headers=flow.headers)).json()["exists"] is False

    edited = copy.deepcopy(ws["plan"])
    edited["objectives"][1]["text"] = "Evaluate ranking quality against three baselines on two datasets."
    edited["title"] = PLAN["title_options"][2]
    saved = await flow.client.put(flow.url("/idea/plan"), json=edited, headers=flow.headers)
    assert saved.status_code == 200, saved.text
    assert saved.json()["plan"]["objectives"][1]["text"].startswith("Evaluate ranking quality against three")

    state = await flow.state()
    assert state["state"]["objectives"][1]["text"].startswith("Evaluate ranking quality against three")
    assert "against three baselines" in state["pending_review"]["draft"]
    title = (await flow.client.get(flow.url("/sections/title"), headers=flow.headers)).json()
    assert title["content"] == PLAN["title_options"][2] and title["current_version"]["source"] == "ai"

    approved = await flow.act("idea", "approve")
    assert next(s for s in approved["stages"] if s["name"] == "idea")["status"] == "approved"
    after = (await flow.client.get(flow.url("/idea"), headers=flow.headers)).json()
    assert after["editable"] is False and after["edit_blocked_reason"] == "Go back to the Idea stage to edit the plan"
    edit_note = select(Feedback.comment).where(
        Feedback.project_id == flow.pid, Feedback.comment == "Edited research plan"
    )
    assert await db_session.scalar(edit_note)


async def test_plan_edits_are_validated(flow: Flow, idea_llm: FakeLLM) -> None:
    idea_llm.responses.append(PLAN)
    await flow.run("idea")
    ws = (await flow.client.get(flow.url("/idea"), headers=flow.headers)).json()

    bad = copy.deepcopy(ws["plan"])
    bad["objectives"] = bad["objectives"][:2] + [{"id": "O4", "text": "Renumbered objective"}]  # RQ2 → O3 now dangling
    resp = await flow.client.put(flow.url("/idea/plan"), json=bad, headers=flow.headers)

    assert resp.status_code == 422
    assert "O3" in str(resp.json()["detail"])


async def test_edited_title_is_marked_as_edited(flow: Flow, idea_llm: FakeLLM) -> None:
    idea_llm.responses.append(PLAN)
    await flow.run("idea")
    ws = (await flow.client.get(flow.url("/idea"), headers=flow.headers)).json()

    edited = {**ws["plan"], "title": "My Own Title for the Paper"}
    assert (await flow.client.put(flow.url("/idea/plan"), json=edited, headers=flow.headers)).status_code == 200

    title = (await flow.client.get(flow.url("/sections/title"), headers=flow.headers)).json()
    assert title["content"] == "My Own Title for the Paper" and title["current_version"]["source"] == "user"


async def test_run_needs_a_raw_idea(flow: Flow, idea_llm: FakeLLM) -> None:
    await flow.client.patch(f"/api/v1/projects/{flow.pid}", json={"raw_idea": ""}, headers=flow.headers)

    resp = await flow.post("/stages/idea/run")

    assert resp.status_code == 409 and resp.json()["detail"] == "Add your research idea to the project first"


async def test_regenerate_from_workspace_uses_feedback(flow: Flow, idea_llm: FakeLLM) -> None:
    idea_llm.responses.extend([PLAN, plan(domain="Medicinal chemistry")])
    await flow.run("idea")

    await flow.act("idea", "regenerate", {"feedback": "Make it about medicinal chemistry"}, expect=202)

    assert "Make it about medicinal chemistry" in idea_llm.calls[1].prompt
    assert "Previous plan" in idea_llm.calls[1].prompt
    ws = (await flow.client.get(flow.url("/idea"), headers=flow.headers)).json()
    assert ws["plan"]["domain"] == "Medicinal chemistry" and ws["status"] == "awaiting_review"


async def test_llm_failure_fails_the_job_and_restores_the_stage(flow: Flow, idea_llm: FakeLLM) -> None:
    idea_llm.responses.append(PermissionError("API key not valid"))
    await flow.post("/stages/idea/run")
    job_id, job_type, params = flow.dispatcher.queued.pop(0)
    from app.db.models.enums import JobStatus
    from app.services.jobs import execute_job
    from app.workers.handlers import HANDLERS

    assert await execute_job(job_id, HANDLERS[job_type], params, factory_for(flow.db)) == JobStatus.FAILED
    job = (await flow.client.get(f"/api/v1/jobs/{job_id}", headers=flow.headers)).json()
    assert "API key not valid" in job["error"]
    assert (await flow.statuses())["idea"] == "not_started"


async def test_viewers_cannot_edit_the_plan(
    flow: Flow, idea_llm: FakeLLM, api_client: AsyncClient, db_session: AsyncSession
) -> None:
    idea_llm.responses.append(PLAN)
    await flow.run("idea")
    bob = await auth_headers(api_client, "bob@example.com")
    bob_id = await db_session.scalar(select(User.id).where(User.email == "bob@example.com"))
    db_session.add(ProjectMember(project_id=flow.pid, user_id=bob_id, role=MemberRole.VIEWER))
    await db_session.flush()

    ws = (await api_client.get(flow.url("/idea"), headers=bob)).json()
    assert ws["plan"] is not None and ws["editable"] is False
    assert (await api_client.put(flow.url("/idea/plan"), json=ws["plan"], headers=bob)).status_code == 403
