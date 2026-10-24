"""Researcher-provided context for agents: text extracted from uploaded files.

Budgets keep prompts bounded: each file is cut to PER_FILE_CHARS and the total to TOTAL_CHARS
(truncation is marked in the text, so the model knows it saw an excerpt).
"""

import asyncio
import io
import json
import uuid
from dataclasses import dataclass

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.logging import get_logger
from app.db.models import File
from app.db.models.enums import FileKind
from app.services.storage import Storage

log = get_logger(__name__)

PER_FILE_CHARS = 12_000
TOTAL_CHARS = 30_000
TEXT_MIMES = {"text/plain", "text/markdown", "application/x-tex", "application/x-bibtex", "text/csv"}
PDF_MIME = "application/pdf"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


@dataclass(frozen=True)
class ContextDocument:
    file_id: uuid.UUID
    filename: str
    text: str
    truncated: bool


def extract_text(data: bytes, mime: str) -> str:
    """Plain text from a file's bytes; empty for formats we don't read."""
    if mime in TEXT_MIMES or mime.startswith("text/"):  # incl. source code (text/x-python, …)
        return data.decode("utf-8", errors="replace")
    if mime == "application/x-ipynb+json":
        return _notebook_text(data)
    if mime == PDF_MIME:
        import pymupdf

        with pymupdf.open(stream=data, filetype="pdf") as doc:  # type: ignore[no-untyped-call]  # untyped library
            return "\n".join(page.get_text() for page in doc)
    if mime == DOCX_MIME:
        import docx

        document = docx.Document(io.BytesIO(data))
        return "\n".join(p.text for p in document.paragraphs)
    return ""


def _notebook_text(data: bytes) -> str:
    """A Jupyter notebook's cells (markdown and code), without outputs."""
    try:
        notebook = json.loads(data)
    except ValueError:
        return ""
    cells = notebook.get("cells") if isinstance(notebook, dict) else None
    parts = []
    for cell in cells or []:
        source = cell.get("source", "")
        parts.append("".join(source) if isinstance(source, list) else str(source))
    return "\n\n".join(p for p in parts if p.strip())


async def load_context(
    session: AsyncSession,
    storage: Storage,
    project_id: uuid.UUID,
    kinds: tuple[FileKind, ...] = (FileKind.NOTES,),
) -> list[ContextDocument]:
    """Text of the project's files of the given kinds, oldest first, within the character budgets."""
    files = await session.scalars(
        select(File).where(File.project_id == project_id, File.kind.in_(kinds)).order_by(File.created_at)
    )
    documents: list[ContextDocument] = []
    remaining = TOTAL_CHARS
    for file in files:
        if remaining <= 0:
            break
        try:
            data = await storage.get_bytes(file.object_key)
            text = await asyncio.to_thread(extract_text, data, file.mime)
        except Exception:
            log.warning("context_file_unreadable", file_id=str(file.id), filename=file.filename, exc_info=True)
            continue
        text = text.strip()
        if not text:
            continue
        limit = min(PER_FILE_CHARS, remaining)
        truncated = len(text) > limit
        documents.append(ContextDocument(file.id, file.filename, text[:limit], truncated))
        remaining -= min(len(text), limit)
    return documents


def format_context(documents: list[ContextDocument]) -> str:
    if not documents:
        return ""
    parts = []
    for doc in documents:
        note = " (excerpt)" if doc.truncated else ""
        parts.append(f'<file name="{doc.filename}"{note}>\n{doc.text}\n</file>')
    return "\n\n".join(parts)
